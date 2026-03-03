#!/usr/bin/env python3
"""
Ingest IMPACT Winter 2026-01 text campaign data from color-coded Excel files.

Creates campaigns, matches/creates participants, and records engagement data
with rich tracking (clicked, form_completed, opted_out).

Color Code Mapping:
  - Blank: Received, no action → unified_status: 'received'
  - Yellow (FFFFFF00): Clicked link, no form → unified_status: 'clicked'
  - Green (FF92D050): Completed form → unified_status: 'converted'
  - Orange (FFFFC000): Negative reply (no stop) → unified_status: 'negative_response'
  - Red (FFFF0000): Replied STOP → unified_status: 'opted_out'
  - Blue (theme_4): Bad number/undeliverable → unified_status: 'undeliverable'

Reads from: data/campaign-IMPACT-202601/*.xlsx
Writes to: empower_development campaigns, participants, exposures

Usage:
    python scripts/ingest_impact_winter_campaign_tool.py [--dry-run]
"""
import argparse
import re
import sys
from datetime import datetime
from collections import defaultdict
from pathlib import Path

import openpyxl
from pymongo import MongoClient

# Campaign configuration
CAMPAIGN_CONFIG = {
    'EmpowerSaves_OH_Text75_Winter_Section1_COLORCODED.xlsx': {
        'campaign_id': 'text_impact_winter_75',
        'name': 'EmpowerSaves OH Text75 Winter Section1 - IMPACT 2026-01',
        'description': 'Winter 2026 IMPACT campaign - Section 1',
    },
    'EmpowerSaves_OH_Text76_Winter_Section2_COLOR CODED.xlsx': {
        'campaign_id': 'text_impact_winter_76',
        'name': 'EmpowerSaves OH Text76 Winter Section2 - IMPACT 2026-01',
        'description': 'Winter 2026 IMPACT campaign - Section 2',
    },
}

# Color code to engagement mapping
COLOR_MAPPING = {
    'FFFFFF00': {  # Yellow - clicked link
        'text_delivered': True, 'text_clicked': True, 'text_form_completed': False,
        'text_replied': False, 'text_opted_out': False, 'unified_status': 'clicked',
    },
    'FF92D050': {  # Green - completed form
        'text_delivered': True, 'text_clicked': True, 'text_form_completed': True,
        'text_replied': False, 'text_opted_out': False, 'unified_status': 'converted',
    },
    'FFFFC000': {  # Orange - negative reply
        'text_delivered': True, 'text_clicked': False, 'text_form_completed': False,
        'text_replied': True, 'text_opted_out': False, 'unified_status': 'negative_response',
    },
    'FFFF0000': {  # Red - opted out (STOP)
        'text_delivered': True, 'text_clicked': False, 'text_form_completed': False,
        'text_replied': True, 'text_opted_out': True, 'unified_status': 'opted_out',
    },
    'theme_4': {  # Blue - undeliverable
        'text_delivered': False, 'text_clicked': False, 'text_form_completed': False,
        'text_replied': False, 'text_opted_out': False, 'unified_status': 'undeliverable',
    },
    None: {  # Blank - received, no action
        'text_delivered': True, 'text_clicked': False, 'text_form_completed': False,
        'text_replied': False, 'text_opted_out': False, 'unified_status': 'received',
    },
}


def normalize_phone(phone) -> str:
    """Normalize phone number to 10 digits."""
    if phone is None:
        return None
    digits = re.sub(r'\D', '', str(phone))
    if len(digits) == 11 and digits.startswith('1'):
        digits = digits[1:]
    return digits if len(digits) == 10 else None


def normalize_email(email) -> str:
    """Normalize email to lowercase."""
    if email is None or str(email).strip() == '':
        return None
    return str(email).lower().strip()


def safe_float(value) -> float:
    """Convert value to float, handling commas, dollar signs, and empty strings."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(',', '').replace('$', '')
    if s == '' or s.lower() == 'nan':
        return None
    try:
        return float(s)
    except ValueError:
        return None


def get_cell_color(cell) -> str:
    """Extract fill color from cell, return color key or None."""
    fill = cell.fill
    if fill.fill_type != 'solid' or not fill.fgColor:
        return None

    color = fill.fgColor
    if color.type == 'rgb' and color.rgb and color.rgb != '00000000':
        return color.rgb
    elif color.type == 'theme':
        return f'theme_{color.theme}'
    return None


def get_engagement_from_color(color_key: str) -> dict:
    """Map color key to engagement fields."""
    return COLOR_MAPPING.get(color_key, COLOR_MAPPING[None])


def read_excel_with_colors(filepath: Path) -> list[dict]:
    """Read Excel file and extract row data with color information."""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active

    # Get headers from first row
    headers = [cell.value for cell in ws[1]]

    rows = []
    for row_idx in range(2, ws.max_row + 1):
        # Get color from first cell
        color_key = get_cell_color(ws.cell(row=row_idx, column=1))

        # Extract row data
        row_data = {'_row_number': row_idx, '_color': color_key}
        for col_idx, header in enumerate(headers, 1):
            if header:  # Skip None headers
                value = ws.cell(row=row_idx, column=col_idx).value
                row_data[header] = value

        rows.append(row_data)

    wb.close()
    return rows


def build_campaign_doc(config: dict) -> dict:
    """Build Campaign document."""
    return {
        'campaign_id': config['campaign_id'],
        'name': config['name'],
        'description': config.get('description', ''),
        'agency': 'IMPACT Community Action',
        'channel': 'text',
        'target_audience': 'IMPACT',
        'campaign_group': 'IMPACT-Winter-202601',
        'message_type': 'SMS',
        'sent_at': datetime(2026, 1, 1),  # Approximate
        'source_system': 'excel_import',
        'source_file': config.get('source_file', ''),
        'statistics': {
            'total_sent': 0,
            'delivered': 0,
            'undeliverable': 0,
            'clicked': 0,
            'converted': 0,
            'negative_response': 0,
            'opted_out': 0,
        },
        'created_at': datetime.utcnow(),
        'synced_at': datetime.utcnow(),
    }


def build_participant_doc(row: dict, phone_norm: str) -> dict:
    """Build new Participant document from Excel row."""
    email = normalize_email(row.get('Email'))

    return {
        'participant_id': phone_norm,
        'email': email,
        'phone': phone_norm,
        'address': {
            'street': str(row.get('ServiceAddress1', '')).strip() or None,
            'city': str(row.get('ServiceCity', '')).strip() or None,
            'zip': str(row.get('ServicePostalCode', ''))[:5] if row.get('ServicePostalCode') else None,
            'county': 'FranklinCounty',  # Default for IMPACT
            'raw': None,
        },
        'linkage': {
            'parcel_id': str(row.get('parcel_id', '')).strip() or None,
            'county_key': 'FranklinCounty',
            'method': 'impact_winter_import',
            'confidence': 0.8 if row.get('parcel_id') else 0.0,
            'matched_at': datetime.utcnow() if row.get('parcel_id') else None,
        },
        'demographics': {},
        'residence': {},
        'energy_snapshot': {},
        'engagement_summary': {
            'total_campaigns': 0,
            'total_exposures': 0,
            'by_channel': {
                'email': {'exposures': 0, 'received': 0, 'engaged': 0},
                'text': {'exposures': 0, 'received': 0, 'engaged': 0},
                'mailer': {'exposures': 0, 'received': 0, 'engaged': 0},
                'letter': {'exposures': 0, 'received': 0, 'engaged': 0},
            },
            'unified_status': 'no_engagement',
            'ever_received': False,
            'ever_engaged': False,
            'first_campaign_date': None,
            'last_campaign_date': None,
        },
        'do_not_contact': False,
        'data_quality': {
            'has_demographics': False,
            'has_residence': False,
            'has_energy_snapshot': False,
            'has_engagement': True,
            'completeness_score': 0.25,
        },
        'created_at': datetime.utcnow(),
        'updated_at': datetime.utcnow(),
        'source_campaigns': [],
    }


def build_exposure_doc(row: dict, participant_id: str, campaign_id: str, engagement: dict) -> dict:
    """Build CampaignExposure document."""
    return {
        'participant_id': participant_id,
        'campaign_id': campaign_id,
        'agency': 'IMPACT Community Action',
        'channel': 'text',
        'sent_at': datetime(2026, 1, 1),  # Approximate send date
        # Email fields (not applicable)
        'email_opened': False,
        'email_clicked': False,
        'email_bounced': False,
        'email_complained': False,
        'email_unsubscribed': False,
        # Text engagement (extended schema)
        'text_delivered': engagement['text_delivered'],
        'text_clicked': engagement['text_clicked'],
        'text_form_completed': engagement['text_form_completed'],
        'text_replied': engagement['text_replied'],
        'text_opted_out': engagement['text_opted_out'],
        'text_read': None,
        # Postal (not applicable)
        'postal_delivered': None,
        'postal_response': None,
        # Unified status
        'unified_status': engagement['unified_status'],
        # Snapshots
        'contact_snapshot': {
            'email': normalize_email(row.get('Email')),
            'phone': normalize_phone(row.get('Cell')),
            'address': str(row.get('ServiceAddress1', '')).strip() or None,
            'city': str(row.get('ServiceCity', '')).strip() or None,
            'zip': str(row.get('ServicePostalCode', ''))[:5] if row.get('ServicePostalCode') else None,
        },
        'energy_at_send': {
            'annual_consumption': safe_float(row.get('annual_consumption')),
            'annual_cost': safe_float(row.get('annual cost')),
            'monthly_cost': safe_float(row.get('monthly cost')),
            'monthly_saving': safe_float(row.get('monthly savings')),
            'daily_cost': safe_float(row.get('daily cost')),
        },
        'created_at': datetime.utcnow(),
    }


def main():
    parser = argparse.ArgumentParser(description='Ingest IMPACT Winter campaign data')
    parser.add_argument('--dry-run', action='store_true', help='Preview without writing to database')
    args = parser.parse_args()

    print("=" * 80)
    print("IMPACT WINTER 2026-01 CAMPAIGN INGESTION")
    print("=" * 80)

    data_dir = Path('/home/yersinia/devel/octopus/data/campaign-IMPACT-202601')

    # Connect to database
    client = MongoClient('localhost', 27017)
    db = client['empower_development']

    # Build existing participant lookup indexes
    print("\n📊 Building participant lookup indexes...")
    existing = list(db.participants.find({}, {
        'participant_id': 1, 'email': 1, 'phone': 1, 'linkage.parcel_id': 1
    }))
    print(f"   Existing participants: {len(existing)}")

    phone_index = {}
    email_index = {}
    parcel_index = {}

    for p in existing:
        pid = p['participant_id']
        if p.get('phone'):
            pn = normalize_phone(p['phone'])
            if pn:
                phone_index[pn] = pid
        if p.get('email'):
            email_index[p['email'].lower()] = pid
        parcel_id = p.get('linkage', {}).get('parcel_id')
        if parcel_id:
            parcel_index[parcel_id] = pid

    print(f"   Phone index: {len(phone_index)} entries")
    print(f"   Email index: {len(email_index)} entries")
    print(f"   Parcel index: {len(parcel_index)} entries")

    # Process each Excel file
    total_stats = {
        'campaigns_created': 0,
        'participants_matched': 0,
        'participants_created': 0,
        'exposures_created': 0,
        'opt_outs_flagged': 0,
    }

    for filename, config in CAMPAIGN_CONFIG.items():
        filepath = data_dir / filename
        if not filepath.exists():
            print(f"\n⚠️  File not found: {filepath}")
            continue

        print(f"\n{'=' * 80}")
        print(f"📁 Processing: {filename}")
        print(f"   Campaign ID: {config['campaign_id']}")
        print("=" * 80)

        # Read Excel with colors
        print("\n📥 Reading Excel file...")
        rows = read_excel_with_colors(filepath)
        print(f"   Loaded {len(rows)} rows")

        # Analyze color distribution
        color_counts = defaultdict(int)
        for row in rows:
            color_counts[row['_color']] += 1

        print("\n🎨 Color distribution:")
        color_labels = {
            None: 'Blank (received)',
            'FFFFFF00': 'Yellow (clicked)',
            'FF92D050': 'Green (converted)',
            'FFFFC000': 'Orange (negative)',
            'FFFF0000': 'Red (opted out)',
            'theme_4': 'Blue (undeliverable)',
        }
        for color, count in sorted(color_counts.items(), key=lambda x: -x[1]):
            label = color_labels.get(color, f'Unknown ({color})')
            print(f"   {label}: {count}")

        if args.dry_run:
            print(f"\n📋 DRY-RUN MODE - No changes will be made")
            continue

        # Create campaign
        print("\n💾 Creating campaign...")
        config['source_file'] = filename
        campaign_doc = build_campaign_doc(config)

        result = db.campaigns.update_one(
            {'campaign_id': config['campaign_id']},
            {'$set': campaign_doc},
            upsert=True
        )
        if result.upserted_id:
            total_stats['campaigns_created'] += 1
            print(f"   Created campaign: {config['campaign_id']}")
        else:
            print(f"   Updated campaign: {config['campaign_id']}")

        # Process participants and exposures
        print("\n👥 Processing participants and exposures...")

        new_participants = []
        exposures = []
        engagement_updates = defaultdict(lambda: {
            'exposures': 0, 'received': 0, 'engaged': 0, 'clicked': 0, 'converted': 0
        })
        opt_out_participants = []

        campaign_stats = defaultdict(int)
        matched = 0
        created = 0

        for row in rows:
            phone_norm = normalize_phone(row.get('Cell'))
            if not phone_norm:
                campaign_stats['skipped_no_phone'] += 1
                continue

            email_norm = normalize_email(row.get('Email'))
            parcel_id = str(row.get('parcel_id', '')).strip() or None

            # Match participant
            participant_id = None

            # Priority 1: Phone match
            if phone_norm in phone_index:
                participant_id = phone_index[phone_norm]
                matched += 1
            # Priority 2: Email match
            elif email_norm and email_norm in email_index:
                participant_id = email_index[email_norm]
                matched += 1
            # Priority 3: Parcel match
            elif parcel_id and parcel_id in parcel_index:
                participant_id = parcel_index[parcel_id]
                matched += 1
            else:
                # Create new participant
                participant_id = phone_norm
                participant_doc = build_participant_doc(row, phone_norm)
                new_participants.append(participant_doc)
                # Add to indexes for dedup within this file
                phone_index[phone_norm] = participant_id
                if email_norm:
                    email_index[email_norm] = participant_id
                if parcel_id:
                    parcel_index[parcel_id] = participant_id
                created += 1

            # Get engagement from color
            color_key = row['_color']
            engagement = get_engagement_from_color(color_key)

            # Build exposure
            exposure = build_exposure_doc(row, participant_id, config['campaign_id'], engagement)
            exposures.append(exposure)

            # Track stats
            status = engagement['unified_status']
            campaign_stats[status] += 1

            # Track engagement updates
            eng = engagement_updates[participant_id]
            eng['exposures'] += 1
            if engagement['text_delivered']:
                eng['received'] += 1
            if engagement['text_clicked']:
                eng['clicked'] += 1
                eng['engaged'] += 1
            if engagement['text_form_completed']:
                eng['converted'] += 1
                eng['engaged'] += 1
            if engagement['text_replied']:
                eng['engaged'] += 1

            # Track opt-outs for global flag
            if engagement['text_opted_out']:
                opt_out_participants.append(participant_id)

        print(f"   Matched existing: {matched}")
        print(f"   New participants: {created}")

        # Insert new participants
        if new_participants:
            print(f"\n📝 Inserting {len(new_participants)} new participants...")
            result = db.participants.insert_many(new_participants)
            total_stats['participants_created'] += len(result.inserted_ids)
            print(f"   Inserted: {len(result.inserted_ids)}")

        total_stats['participants_matched'] += matched

        # Insert exposures
        print(f"\n📝 Inserting {len(exposures)} exposures...")
        if exposures:
            # Use ordered=False to continue on duplicates
            try:
                result = db.exposures.insert_many(exposures, ordered=False)
                total_stats['exposures_created'] += len(result.inserted_ids)
                print(f"   Inserted: {len(result.inserted_ids)}")
            except Exception as e:
                if 'duplicate' in str(e).lower():
                    print(f"   Some duplicates skipped")
                else:
                    print(f"   ⚠️ Error: {e}")

        # Update campaign statistics
        print("\n📊 Updating campaign statistics...")
        db.campaigns.update_one(
            {'campaign_id': config['campaign_id']},
            {'$set': {
                'statistics.total_sent': len(rows),
                'statistics.delivered': campaign_stats.get('received', 0) +
                                        campaign_stats.get('clicked', 0) +
                                        campaign_stats.get('converted', 0) +
                                        campaign_stats.get('negative_response', 0) +
                                        campaign_stats.get('opted_out', 0),
                'statistics.undeliverable': campaign_stats.get('undeliverable', 0),
                'statistics.clicked': campaign_stats.get('clicked', 0),
                'statistics.converted': campaign_stats.get('converted', 0),
                'statistics.negative_response': campaign_stats.get('negative_response', 0),
                'statistics.opted_out': campaign_stats.get('opted_out', 0),
            }}
        )

        # Update participant engagement summaries
        print(f"\n📊 Updating {len(engagement_updates)} participant engagement summaries...")
        for participant_id, eng in engagement_updates.items():
            update = {
                '$inc': {
                    'engagement_summary.total_exposures': eng['exposures'],
                    'engagement_summary.by_channel.text.exposures': eng['exposures'],
                    'engagement_summary.by_channel.text.received': eng['received'],
                    'engagement_summary.by_channel.text.engaged': eng['engaged'],
                },
                '$addToSet': {'source_campaigns': config['campaign_id']},
                '$set': {
                    'updated_at': datetime.utcnow(),
                    'data_quality.has_engagement': True,
                }
            }
            db.participants.update_one({'participant_id': participant_id}, update)

        # Flag opt-outs globally
        if opt_out_participants:
            print(f"\n🚫 Flagging {len(opt_out_participants)} opt-outs (do_not_contact)...")
            result = db.participants.update_many(
                {'participant_id': {'$in': opt_out_participants}},
                {'$set': {
                    'do_not_contact': True,
                    'engagement_summary.unified_status': 'opted_out',
                    'updated_at': datetime.utcnow(),
                }}
            )
            total_stats['opt_outs_flagged'] += result.modified_count
            print(f"   Flagged: {result.modified_count}")

        print(f"\n✅ Completed: {config['campaign_id']}")
        print(f"   Status breakdown: {dict(campaign_stats)}")

    # Final summary
    print("\n" + "=" * 80)
    print("📊 INGESTION SUMMARY")
    print("=" * 80)

    if args.dry_run:
        print("\n  DRY-RUN MODE - No changes were made")
        print("  Run without --dry-run to execute ingestion")
    else:
        print(f"\n  Campaigns created:     {total_stats['campaigns_created']}")
        print(f"  Participants matched:  {total_stats['participants_matched']}")
        print(f"  Participants created:  {total_stats['participants_created']}")
        print(f"  Exposures created:     {total_stats['exposures_created']}")
        print(f"  Opt-outs flagged:      {total_stats['opt_outs_flagged']}")

        # Verify totals
        print("\n  Database totals:")
        print(f"    Total campaigns:  {db.campaigns.count_documents({})}")
        print(f"    Total participants: {db.participants.count_documents({})}")
        print(f"    Total exposures:    {db.exposures.count_documents({})}")

    client.close()
    print("\n✅ Done!")

    return 0


if __name__ == '__main__':
    sys.exit(main())
